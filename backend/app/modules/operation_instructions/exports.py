import io
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from xml.sax.saxutils import escape

import qrcode
import qrcode.image.svg
from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph as DocxParagraph
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.core.config import get_settings
from app.modules.operation_instructions.markdown import markdown_to_text

IMAGE_PATTERN = re.compile(r"!\[([^]]*)]\(([^)\s]+)(?:\s+['\"][^'\"]*['\"])?\)")
TOKEN_PATTERN = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*|`[^`]+`)")


@dataclass(frozen=True, slots=True)
class ExportedInstruction:
    body: bytes
    media_type: str
    filename: str


def _safe_name(name: str) -> str:
    cleaned = re.sub(r"[^\w.-]+", "-", name, flags=re.UNICODE).strip("-.")
    return cleaned[:80] or "instruction"


def _local_image(source: str) -> Path | None:
    filename = source.rsplit("/media/", 1)[-1].split("?", 1)[0]
    if not filename or "/" in filename or "\\" in filename:
        return None
    path = (get_settings().media_root / filename).resolve()
    root = get_settings().media_root.resolve()
    return path if path.parent == root and path.is_file() else None


def _docx_inline(paragraph: DocxParagraph, text: str) -> None:
    add_run = paragraph.add_run
    cursor = 0
    for match in TOKEN_PATTERN.finditer(text):
        add_run(text[cursor : match.start()])
        token = match.group(0)
        run = add_run(token[2:-2] if token.startswith("**") else token[1:-1])
        if token.startswith("**"):
            run.bold = True
        elif token.startswith("*"):
            run.italic = True
        else:
            run.font.name = "Consolas"
        cursor = match.end()
    add_run(text[cursor:])


def _docx(content: str, operation_name: str, version_number: int) -> bytes:
    document = Document()
    document.add_heading(operation_name, level=0)
    document.add_paragraph(f"Инструкция · версия {version_number}")
    lines = content.splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index]
        if line.strip().startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            run = document.add_paragraph().add_run(line)
            run.font.name = "Consolas"
            index += 1
            continue
        image_match = IMAGE_PATTERN.fullmatch(line.strip())
        if image_match:
            path = _local_image(image_match.group(2))
            if path:
                document.add_picture(str(path), width=Inches(6.0))
            if image_match.group(1):
                document.add_paragraph(image_match.group(1), style="Caption")
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and "---" in lines[index + 1]:
            rows: list[list[str]] = []
            rows.append([cell.strip() for cell in line.strip("|").split("|")])
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip("|").split("|")])
                index += 1
            table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
            table.style = "Table Grid"
            for row_index, cells in enumerate(rows):
                for cell_index, value in enumerate(cells):
                    _docx_inline(table.cell(row_index, cell_index).paragraphs[0], value)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
        elif re.match(r"^\s*[-*+]\s+", line):
            _docx_inline(
                document.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*+]\s+", "", line)
            )
        elif re.match(r"^\s*\d+[.)]\s+", line):
            _docx_inline(
                document.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line)
            )
        elif line.startswith(">"):
            _docx_inline(document.add_paragraph(style="Quote"), line.lstrip("> "))
        elif line.strip() in {"---", "***", "___"}:
            document.add_paragraph("―" * 32)
        elif line.strip():
            _docx_inline(document.add_paragraph(), line)
        index += 1
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _font_name() -> str:
    candidates = (
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    for path in candidates:
        if path.is_file():
            name = "WebStorageUnicode"
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(path)))
            return name
    return "Helvetica"


def _pdf(content: str, operation_name: str, version_number: int) -> bytes:
    output = io.BytesIO()
    font = _font_name()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=operation_name,
    )
    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "Instruction",
        parent=styles["BodyText"],
        fontName=font,
        fontSize=10,
        leading=14,
        spaceAfter=6,
    )
    heading = ParagraphStyle(
        "InstructionTitle",
        parent=normal,
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
        spaceAfter=10,
    )
    story: list[object] = [Paragraph(escape(operation_name), heading)]
    story.append(Paragraph(f"Инструкция · версия {version_number}", normal))
    story.append(Paragraph(f"Сформировано: {datetime.now(UTC):%d.%m.%Y %H:%M UTC}", normal))
    story.append(Spacer(1, 5 * mm))
    for line in content.splitlines():
        stripped = line.strip()
        image_match = IMAGE_PATTERN.fullmatch(stripped)
        if image_match:
            path = _local_image(image_match.group(2))
            if path:
                picture = Image(str(path))
                picture._restrictSize(170 * mm, 110 * mm)
                story.append(picture)
            if image_match.group(1):
                story.append(Paragraph(escape(image_match.group(1)), normal))
            continue
        title_match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if title_match:
            size = max(11, 18 - len(title_match.group(1)) * 2)
            story.append(
                Paragraph(
                    escape(title_match.group(2)),
                    ParagraphStyle(
                        f"H{size}",
                        parent=normal,
                        fontSize=size,
                        leading=size + 4,
                        spaceBefore=8,
                        spaceAfter=5,
                    ),
                )
            )
        elif re.match(r"^\s*[-*+]\s+", line):
            value = re.sub(r"^\s*[-*+]\s+", "", line)
            story.append(
                ListFlowable([ListItem(Paragraph(escape(value), normal))], bulletType="bullet")
            )
        elif stripped and stripped not in {"---", "***", "___"}:
            story.append(Paragraph(escape(stripped.lstrip("> ")), normal))
        elif stripped:
            story.append(
                Table(
                    [[""]],
                    colWidths=[170 * mm],
                    style=TableStyle([("LINEABOVE", (0, 0), (-1, -1), 0.5, colors.grey)]),
                )
            )
    document.build(story)
    return output.getvalue()


def export_instruction(
    export_format: str,
    *,
    operation_name: str,
    content: str,
    version_number: int,
) -> ExportedInstruction:
    basename = f"{_safe_name(operation_name)}-v{version_number}"
    if export_format == "md":
        return ExportedInstruction(
            content.encode(), "text/markdown; charset=utf-8", f"{basename}.md"
        )
    if export_format == "txt":
        return ExportedInstruction(
            markdown_to_text(content).encode(), "text/plain; charset=utf-8", f"{basename}.txt"
        )
    if export_format == "docx":
        return ExportedInstruction(
            _docx(content, operation_name, version_number),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{basename}.docx",
        )
    if export_format == "pdf":
        return ExportedInstruction(
            _pdf(content, operation_name, version_number), "application/pdf", f"{basename}.pdf"
        )
    raise ValueError("Unsupported export format")


def qr_svg(public_url: str) -> bytes:
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, border=3)
    qr.add_data(public_url)
    qr.make(fit=True)
    image = qr.make_image(image_factory=qrcode.image.svg.SvgPathImage)
    output = io.BytesIO()
    image.save(output)
    return output.getvalue()
